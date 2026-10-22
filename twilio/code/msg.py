import os
import json
import time
import logging
import csv
from openpyxl import load_workbook
from dotenv import load_dotenv
from twilio.rest import Client
from twilio.base.exceptions import TwilioRestException

# 1. Setup Logging
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)

# 2. Load Credentials securely from .env
load_dotenv()
account_sid = os.environ.get("TWILIO_ACCOUNT_SID")
auth_token = os.environ.get("TWILIO_AUTH_TOKEN")

if not account_sid or not auth_token:
    raise SystemExit("ERROR: TWILIO_ACCOUNT_SID or TWILIO_AUTH_TOKEN missing in .env file.")

client = Client(account_sid, auth_token)

# 3. Configuration
FROM_NUMBER  = "whatsapp:+14155238886"                # Twilio Sandbox number
CONTENT_SID  = "HXxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx"    # PASTE YOUR NEW APPROVED TEMPLATE SID HERE
XLSX_FILE    = "contacts.xlsx"
DELAY_SEC    = 1.0                                    # Delay to avoid rate limits

# 4. Read Excel File
try:
    wb = load_workbook(XLSX_FILE, read_only=True, data_only=True)
    ws = wb.active
except FileNotFoundError:
    raise SystemExit(f"ERROR: Could not find {XLSX_FILE}. Make sure it is in the same folder.")

headers = [str(c.value).strip() if c.value else "" for c in ws[1]]

if "phone" not in headers or "image" not in headers:
    raise SystemExit("ERROR: Excel must have 'phone' and 'image' headers in row 1.")

phone_idx = headers.index("phone")
image_idx = headers.index("image")

seen = set()
sent, failed = 0, 0
results = []

log.info("Starting bulk send process...")

# 5. Loop through rows
for row in ws.iter_rows(min_row=2, values_only=True):
    phone = row[phone_idx]
    if not phone:
        continue
    
    phone = str(phone).strip()
    
    # Validate phone number format
    if not phone.startswith("+"):
        log.warning("Skipping %s — missing '+' country code", phone)
        failed += 1
        results.append((phone, "", "invalid_phone_format"))
        continue
        
    if phone in seen:
        log.info("Skipping duplicate number: %s", phone)
        continue
    seen.add(phone)

    # Validate image
    image_name = str(row[image_idx]).strip() if row[image_idx] else ""
    if not image_name:
        log.warning("Skipping %s — missing image filename", phone)
        failed += 1
        results.append((phone, "", "missing_image"))
        continue

    # 6. Build Variables (Mapping {{1}} to the image)
    variables = {
        "1": image_name
    }

    # 7. Send Message
    try:
        msg = client.messages.create(
            from_=FROM_NUMBER,
            to=f"whatsapp:{phone}",
            content_sid=CONTENT_SID,
            content_variables=json.dumps(variables)
        )
        log.info("Sent -> %s | Image: %s | SID: %s", phone, image_name, msg.sid)
        sent += 1
        results.append((phone, msg.sid, msg.status))
        
    except TwilioRestException as e:
        log.error("Failed -> %s | Error Code: %s | Msg: %s", phone, e.code, e.msg)
        failed += 1
        results.append((phone, "", f"error_{e.code}"))

    # 8. Delay
    time.sleep(DELAY_SEC)

wb.close()

# 9. Generate Report
report_file = "send_report.csv"
with open(report_file, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(["phone", "sid", "status"])
    writer.writerows(results)

log.info("Finished! Sent: %d, Failed: %d. Check %s for details.", sent, failed, report_file)